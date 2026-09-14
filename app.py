import os
from datetime import datetime

from flask import (
    Flask,
    render_template,
    request,
    redirect,
    url_for,
    flash,
    send_from_directory,
    send_file
)

from io import BytesIO

from flask_login import (
    LoginManager,
    login_user,
    logout_user,
    login_required,
    current_user
)

from typing import cast

from sqlalchemy import text

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.cell.cell import Cell

from werkzeug.utils import secure_filename

from config import Config

from material_request import (
    db,
    MaterialRequest,
    User,
    ApprovalHistory,
    Notification
)


# =========================================================
# APPLICATION
# =========================================================

app = Flask(__name__)

app.config.from_object(Config)


# =========================================================
# CONFIGURATION
# =========================================================

UPLOAD_FOLDER = os.path.join(
    app.root_path,
    "uploads"
)

app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER

ALLOWED_EXTENSIONS = {
    "xlsx",
    "xls",
    "pdf"
}

os.makedirs(
    UPLOAD_FOLDER,
    exist_ok=True
)


# =========================================================
# DATABASE
# =========================================================

db.init_app(app)


# =========================================================
# LOGIN MANAGER
# =========================================================

login_manager = LoginManager()

login_manager.init_app(app)

setattr(
    login_manager,
    "login_view",
    "login"
)

login_manager.login_message = (
    "Please log in to access this page."
)


@login_manager.user_loader
def load_user(user_id):

    return User.query.get(
        int(user_id)
    )


# =========================================================
# HELPER FUNCTIONS
# =========================================================

def allowed_file(
    filename,
    extensions=None
):

    if "." not in filename:
        return False

    extension = filename.rsplit(
        ".",
        1
    )[1].lower()

    if extensions is None:
        extensions = ALLOWED_EXTENSIONS

    return extension in extensions


# =========================================================
# NOTIFICATION HELPER
# =========================================================

def create_notification(
    user_id,
    title,
    message,
    notification_type="INFO",
    material_request_id=None
):

    notification = Notification()

    notification.user_id = user_id

    notification.title = title

    notification.message = message

    notification.notification_type = (
        notification_type
    )

    notification.material_request_id = (
        material_request_id
    )

    db.session.add(
        notification
    )


# =========================================================
# NOTIFY USERS BY ROLE
# =========================================================

def notify_users_by_role(
    role,
    title,
    message,
    notification_type="INFO",
    material_request_id=None
):

    users = User.query.filter_by(
        role=role,
        is_active=True
    ).all()

    for user in users:

        create_notification(
            user_id=user.id,
            title=title,
            message=message,
            notification_type=notification_type,
            material_request_id=material_request_id
        )


# =========================================================
# NOTIFY REQUESTER
# =========================================================

def notify_requester(
    material_request,
    title,
    message,
    notification_type="INFO"
):

    requester = User.query.filter_by(
        full_name=material_request.requested_by,
        is_active=True
    ).first()

    if requester:

        create_notification(
            user_id=requester.id,
            title=title,
            message=message,
            notification_type=notification_type,
            material_request_id=material_request.id
        )


# =========================================================
# NOTIFY ADMINS
# =========================================================

def notify_admins(
    title,
    message,
    notification_type="INFO",
    material_request_id=None
):

    notify_users_by_role(
        role="ADMIN",
        title=title,
        message=message,
        notification_type=notification_type,
        material_request_id=material_request_id
    )


# =========================================================
# LOGIN
# =========================================================

@app.route(
    "/login",
    methods=["GET", "POST"]
)
def login():

    if current_user.is_authenticated:

        return redirect(
            url_for("home")
        )

    if request.method == "POST":

        username = request.form.get(
            "username",
            ""
        ).strip()

        password = request.form.get(
            "password",
            ""
        )

        user = User.query.filter_by(
            username=username
        ).first()

        if user and user.check_password(
            password
        ):

            if not user.is_active:

                flash(
                    "Your account is inactive.",
                    "danger"
                )

                return redirect(
                    url_for("login")
                )

            login_user(user)

            return redirect(
                url_for("home")
            )

        flash(
            "Invalid username or password.",
            "danger"
        )

    return render_template(
        "login.html"
    )


# =========================================================
# LOGOUT
# =========================================================

@app.route("/logout")
@login_required
def logout():

    logout_user()

    flash(
        "You have been logged out.",
        "success"
    )

    return redirect(
        url_for("login")
    )


# =========================================================
# DASHBOARD
# =========================================================

@app.route("/")
@login_required
def home():

    search = request.args.get(
        "search",
        ""
    ).strip()

    status = request.args.get(
        "status",
        ""
    ).strip()

    query = MaterialRequest.query

    # -----------------------------------------------------
    # SEARCH
    # -----------------------------------------------------

    if search:

        search_filter = (
            MaterialRequest.mri_no.ilike(
                f"%{search}%"
            )
            |
            MaterialRequest.company.ilike(
                f"%{search}%"
            )
            |
            MaterialRequest.requested_by.ilike(
                f"%{search}%"
            )
        )

        query = query.filter(
            search_filter
        )

    # -----------------------------------------------------
    # STATUS FILTER
    # -----------------------------------------------------

    if status:

        query = query.filter(
            MaterialRequest.material_status
            == status
        )

    requests = query.order_by(
        MaterialRequest.id.desc()
    ).all()

    # -----------------------------------------------------
    # DASHBOARD COUNTS
    # -----------------------------------------------------

    total = MaterialRequest.query.count()

    piping_approved = (
        MaterialRequest.query
        .filter_by(
            piping_status="Approved"
        )
        .count()
    )

    piping_rejected = (
        MaterialRequest.query
        .filter_by(
            piping_status="Rejected"
        )
        .count()
    )

    material_approved = (
        MaterialRequest.query
        .filter_by(
            material_status="Approved"
        )
        .count()
    )

    material_rejected = (
        MaterialRequest.query
        .filter_by(
            material_status="Rejected"
        )
        .count()
    )

    return render_template(
        "index.html",
        requests=requests,
        search=search,
        status=status,
        total=total,
        piping_approved=piping_approved,
        piping_rejected=piping_rejected,
        material_approved=material_approved,
        material_rejected=material_rejected
    )


# =========================================================
# EXPORT MATERIAL REQUESTS TO EXCEL
# =========================================================

@app.route("/export/excel")
@login_required
def export_excel():

    search = request.args.get(
        "search",
        ""
    ).strip()

    status = request.args.get(
        "status",
        ""
    ).strip()

    query = MaterialRequest.query

    if search:

        search_filter = (
            MaterialRequest.mri_no.ilike(
                f"%{search}%"
            )
            |
            MaterialRequest.company.ilike(
                f"%{search}%"
            )
            |
            MaterialRequest.requested_by.ilike(
                f"%{search}%"
            )
        )

        query = query.filter(
            search_filter
        )

    if status:

        query = query.filter(
            MaterialRequest.material_status
            == status
        )

    requests = query.order_by(
        MaterialRequest.id.desc()
    ).all()

    # -----------------------------------------------------
    # CREATE WORKBOOK
    # -----------------------------------------------------

    workbook = Workbook()

    worksheet = cast(
        Worksheet,
        workbook.active
    )

    worksheet.title = (
        "Material Requests"
    )

    # -----------------------------------------------------
    # TITLE
    # -----------------------------------------------------

    worksheet["A1"] = (
        "MATERIAL REQUEST REPORT"
    )

    worksheet["A1"].font = Font(
        bold=True,
        size=16
    )

    worksheet["A2"] = (
        "Generated: "
        f"{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    )

    worksheet["A3"] = (
        f"Total Records: {len(requests)}"
    )

    # -----------------------------------------------------
    # FILTERS
    # -----------------------------------------------------

    worksheet["A4"] = "Search:"
    worksheet["B4"] = (
        search if search else "All"
    )

    worksheet["C4"] = (
        "Material Status:"
    )

    worksheet["D4"] = (
        status if status else "All"
    )

    # -----------------------------------------------------
    # HEADERS
    # -----------------------------------------------------

    headers = [
        "ID",
        "Company",
        "MRI No.",
        "Requested By",
        "Request Date",
        "LTSC Status",
        "Piping Status",
        "Piping Date",
        "CCS-JV Confirmation",
        "Material Status",
        "Material Date",
        "Excel File",
        "PDF File",
        "Voucher Files",
        "Remarks",
        "Created At",
        "Updated At"
    ]

    header_row = 6

    for column, header in enumerate(
        headers,
        start=1
    ):

        cell = cast(
            Cell,
            worksheet.cell(
                row=header_row,
                column=column
            )
        )

        cell.value = header

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    # -----------------------------------------------------
    # DATA
    # -----------------------------------------------------

    row_number = header_row + 1

    for item in requests:

        values = [
            item.id,
            item.company,
            item.mri_no,
            item.requested_by,
            item.req_date,
            item.ltsc_status,
            item.piping_status,
            item.piping_date,
            item.ccs_jv_confirm,
            item.material_status,
            item.mtrl_date,
            item.attachment_excel,
            item.attachment_pdf,
            item.voucher_files,
            item.remarks,
            item.created_at,
            item.updated_at
        ]

        for column, value in enumerate(
            values,
            start=1
        ):

            cell = cast(
                Cell,
                worksheet.cell(
                    row=row_number,
                    column=column
                )
            )

            cell.value = value

            cell.alignment = Alignment(
                vertical="top"
            )

        worksheet.cell(
            row=row_number,
            column=5
        ).number_format = (
            "yyyy-mm-dd"
        )

        worksheet.cell(
            row=row_number,
            column=8
        ).number_format = (
            "yyyy-mm-dd"
        )

        worksheet.cell(
            row=row_number,
            column=11
        ).number_format = (
            "yyyy-mm-dd"
        )

        worksheet.cell(
            row=row_number,
            column=16
        ).number_format = (
            "yyyy-mm-dd hh:mm:ss"
        )

        worksheet.cell(
            row=row_number,
            column=17
        ).number_format = (
            "yyyy-mm-dd hh:mm:ss"
        )

        row_number += 1

    # -----------------------------------------------------
    # FREEZE
    # -----------------------------------------------------

    worksheet.freeze_panes = "A7"

    # -----------------------------------------------------
    # FILTER
    # -----------------------------------------------------

    if requests:

        last_row = (
            header_row
            + len(requests)
        )

        worksheet.auto_filter.ref = (
            f"A{header_row}:Q{last_row}"
        )

    # -----------------------------------------------------
    # WIDTHS
    # -----------------------------------------------------

    widths = {
        "A": 10,
        "B": 20,
        "C": 18,
        "D": 22,
        "E": 15,
        "F": 15,
        "G": 18,
        "H": 15,
        "I": 22,
        "J": 20,
        "K": 15,
        "L": 30,
        "M": 30,
        "N": 40,
        "O": 40,
        "P": 22,
        "Q": 22
    }

    for column, width in widths.items():

        worksheet.column_dimensions[
            column
        ].width = width

    worksheet.row_dimensions[
        header_row
    ].height = 30

    # -----------------------------------------------------
    # FILE
    # -----------------------------------------------------

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    filename = (
        "Material_Request_Report_"
        f"{datetime.now().strftime('%Y-%m-%d')}.xlsx"
    )

    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )


# =========================================================
# NEW MATERIAL REQUEST
# =========================================================

@app.route(
    "/material-request/new",
    methods=["GET", "POST"]
)
@login_required
def new_material_request():

    # =====================================================
    # ONLY ADMIN AND REQUESTER CAN CREATE
    # =====================================================

    allowed_roles = [
        "ADMIN",
        "REQUESTER"
    ]

    if current_user.role not in allowed_roles:

        flash(
            "You do not have permission to create "
            "a Material Request.",
            "danger"
        )

        return redirect(
            url_for("home")
        )

    if request.method == "POST":

        # =================================================
        # BASIC INFORMATION
        # =================================================

        company = request.form.get(
            "company",
            ""
        ).strip()

        mri_no = request.form.get(
            "mri_no",
            ""
        ).strip()

        requested_by = request.form.get(
            "requested_by",
            ""
        ).strip()

        req_date = request.form.get(
            "req_date",
            ""
        ).strip()

        # -------------------------------------------------
        # VALIDATION
        # -------------------------------------------------

        if not company:

            flash(
                "Company is required.",
                "danger"
            )

            return redirect(
                url_for(
                    "new_material_request"
                )
            )

        if not mri_no:

            flash(
                "MRI Number is required.",
                "danger"
            )

            return redirect(
                url_for(
                    "new_material_request"
                )
            )

        if not requested_by:

            flash(
                "Requested By is required.",
                "danger"
            )

            return redirect(
                url_for(
                    "new_material_request"
                )
            )

        if not req_date:

            flash(
                "Request Date is required.",
                "danger"
            )

            return redirect(
                url_for(
                    "new_material_request"
                )
            )

        try:

            req_date_obj = (
                datetime.strptime(
                    req_date,
                    "%Y-%m-%d"
                ).date()
            )

        except ValueError:

            flash(
                "Invalid Request Date.",
                "danger"
            )

            return redirect(
                url_for(
                    "new_material_request"
                )
            )

        # =================================================
        # DEFAULT VALUES
        # =================================================

        ltsc_status = "Submitted"

        piping_status = "Pending"
        piping_date = None

        ccs_jv_confirm = "Pending"

        material_status = "Pending"
        mtrl_date = None

        remarks = request.form.get(
            "remarks",
            ""
        ).strip()

        # =================================================
        # ADMIN
        # =================================================

        if current_user.role == "ADMIN":

            ltsc_status = request.form.get(
                "ltsc_status",
                "Submitted"
            )

            if ltsc_status not in [
                "Submitted",
                "Cancelled"
            ]:

                ltsc_status = "Submitted"

            piping_status = request.form.get(
                "piping_status",
                "Pending"
            )

            if piping_status not in [
                "Pending",
                "In Progress",
                "Completed",
                "Approved",
                "Rejected"
            ]:

                piping_status = "Pending"

            piping_date_value = (
                request.form.get(
                    "piping_date",
                    ""
                ).strip()
            )

            if piping_date_value:

                try:

                    piping_date = (
                        datetime.strptime(
                            piping_date_value,
                            "%Y-%m-%d"
                        ).date()
                    )

                except ValueError:

                    flash(
                        "Invalid Piping Date.",
                        "danger"
                    )

                    return redirect(
                        url_for(
                            "new_material_request"
                        )
                    )

            ccs_jv_confirm = request.form.get(
                "ccs_jv_confirm",
                "Pending"
            )

            if ccs_jv_confirm not in [
                "Pending",
                "Confirmed",
                "Rejected"
            ]:

                ccs_jv_confirm = "Pending"

            material_status = request.form.get(
                "material_status",
                "Pending"
            )

            if material_status not in [
                "Pending",
                "Approved",
                "Ordered",
                "In Transit",
                "Received",
                "Rejected"
            ]:

                material_status = "Pending"

            mtrl_date_value = (
                request.form.get(
                    "mtrl_date",
                    ""
                ).strip()
            )

            if mtrl_date_value:

                try:

                    mtrl_date = (
                        datetime.strptime(
                            mtrl_date_value,
                            "%Y-%m-%d"
                        ).date()
                    )

                except ValueError:

                    flash(
                        "Invalid Material Date.",
                        "danger"
                    )

                    return redirect(
                        url_for(
                            "new_material_request"
                        )
                    )

        # =================================================
        # REQUESTER
        # =================================================

        elif current_user.role == "REQUESTER":

            ltsc_status = request.form.get(
                "ltsc_status",
                "Submitted"
            )

            if ltsc_status not in [
                "Submitted",
                "Cancelled"
            ]:

                ltsc_status = "Submitted"

        # =================================================
        # FILES
        # =================================================

        excel_file = request.files.get(
            "attachment_excel"
        )

        pdf_file = request.files.get(
            "attachment_pdf"
        )

        voucher_files = request.files.getlist(
            "voucher_files"
        )

        excel_filename = None

        pdf_filename = None

        voucher_filenames = []

        # =================================================
        # EXCEL
        # =================================================

        if (
            excel_file
            and excel_file.filename
        ):

            if not allowed_file(
                excel_file.filename,
                {"xls", "xlsx"}
            ):

                flash(
                    "Invalid Excel file.",
                    "danger"
                )

                return redirect(
                    url_for(
                        "new_material_request"
                    )
                )

            excel_filename = secure_filename(
                excel_file.filename
            )

            excel_file.save(
                os.path.join(
                    app.config["UPLOAD_FOLDER"],
                    excel_filename
                )
            )

        # =================================================
        # PDF
        # =================================================

        if (
            pdf_file
            and pdf_file.filename
        ):

            if not allowed_file(
                pdf_file.filename,
                {"pdf"}
            ):

                flash(
                    "Invalid PDF file.",
                    "danger"
                )

                return redirect(
                    url_for(
                        "new_material_request"
                    )
                )

            pdf_filename = secure_filename(
                pdf_file.filename
            )

            pdf_file.save(
                os.path.join(
                    app.config["UPLOAD_FOLDER"],
                    pdf_filename
                )
            )

        # =================================================
        # VOUCHER FILES
        # ADMIN ONLY
        # =================================================

        if current_user.role == "ADMIN":

            for voucher_file in voucher_files:

                if (
                    not voucher_file
                    or not voucher_file.filename
                ):
                    continue

                if not allowed_file(
                    voucher_file.filename,
                    {"pdf"}
                ):

                    flash(
                        "Voucher files must be PDF.",
                        "danger"
                    )

                    return redirect(
                        url_for(
                            "new_material_request"
                        )
                    )

                voucher_filename = secure_filename(
                    voucher_file.filename
                )

                voucher_file.save(
                    os.path.join(
                        app.config["UPLOAD_FOLDER"],
                        voucher_filename
                    )
                )

                voucher_filenames.append(
                    voucher_filename
                )
                
            # =================================================
            # CHECK DUPLICATE MRI NUMBER
            # =================================================

            existing_mri = MaterialRequest.query.filter_by(
            mri_no=mri_no
            ).first()

            if existing_mri:

                flash(
                f"ERROR: MRI Number '{mri_no}' already exists. "
                "Please enter a different MRI Number.",
                "danger"
            )

            return redirect(
                url_for("new_material_request")
            )

        # =================================================
        # CREATE MATERIAL REQUEST
        # =================================================

        material_request = MaterialRequest()

        material_request.company = company

        material_request.mri_no = mri_no

        material_request.requested_by = (
            requested_by
        )

        material_request.req_date = (
            req_date_obj
        )

        material_request.ltsc_status = (
            ltsc_status
        )

        material_request.piping_status = (
            piping_status
        )

        material_request.piping_date = (
            piping_date
        )

        material_request.ccs_jv_confirm = (
            ccs_jv_confirm
        )

        material_request.material_status = (
            material_status
        )

        material_request.mtrl_date = (
            mtrl_date
        )

        material_request.attachment_excel = (
            excel_filename
        )

        material_request.attachment_pdf = (
            pdf_filename
        )

        material_request.voucher_files = (
            ",".join(voucher_filenames)
        )

        material_request.remarks = (
            remarks
        )

        # =================================================
        # SAVE REQUEST + NOTIFICATIONS
        # =================================================

        try:

            db.session.add(
                material_request
            )

            # Get ID before creating notifications
            db.session.flush()

            # =================================================
            # NOTIFY PIPING APPROVERS
            # =================================================

            notify_users_by_role(
                role="APPROVER_PIPING",
                title="New Material Request",
                message=(
                    f"{material_request.mri_no} "
                    f"has been submitted and is waiting "
                    f"for Piping Approval."
                ),
                notification_type="NEW_REQUEST",
                material_request_id=(
                    material_request.id
                )
            )

            # =================================================
            # NOTIFY ADMINS
            # =================================================

            notify_admins(
                title="New Material Request",
                message=(
                    f"{material_request.mri_no} "
                    f"has been submitted by "
                    f"{material_request.requested_by}."
                ),
                notification_type="NEW_REQUEST",
                material_request_id=(
                    material_request.id
                )
            )

            # =================================================
            # COMMIT
            # =================================================

            db.session.commit()

            flash(
                "Material Request created successfully.",
                "success"
            )

            return redirect(
                url_for("home")
            )

        except Exception as e:

            db.session.rollback()

            flash(
                f"Error creating request: {e}",
                "danger"
            )

            return redirect(
                url_for(
                    "new_material_request"
                )
            )

    return render_template(
        "material_request_form.html"
    )


# =========================================================
# VIEW MATERIAL REQUEST
# =========================================================

@app.route(
    "/material-request/<int:id>"
)
@login_required
def view_material_request(id):

    material_request = (
        MaterialRequest.query
        .get_or_404(id)
    )

    return render_template(
        "material_request_view.html",
        item=material_request
    )


# =========================================================
# EDIT MATERIAL REQUEST
# =========================================================

@app.route(
    "/material-request/<int:id>/edit",
    methods=["GET", "POST"]
)
@login_required
def edit_material_request(id):

    material_request = (
        MaterialRequest.query
        .get_or_404(id)
    )

    role = current_user.role

    # =====================================================
    # REMEMBER ORIGINAL STATUS
    # =====================================================

    original_ltsc_status = (
        material_request.ltsc_status
    )

    original_piping_status = (
        material_request.piping_status
    )

    original_material_status = (
        material_request.material_status
    )

    # =====================================================
    # ALLOWED ROLES
    # =====================================================

    allowed_roles = [
        "ADMIN",
        "REQUESTER",
        "APPROVER_PIPING",
        "APPROVER_MATERIAL"
    ]

    if role not in allowed_roles:

        flash(
            "You do not have permission to edit "
            "this request.",
            "danger"
        )

        return redirect(
            url_for(
                "view_material_request",
                id=id
            )
        )

    # =====================================================
    # REQUESTER ACCESS
    # =====================================================

    if role == "REQUESTER":

        if (
            material_request.ltsc_status
            != "Submitted"
        ):

            flash(
                "This request can no longer be "
                "edited by the requester.",
                "danger"
            )

            return redirect(
                url_for(
                    "view_material_request",
                    id=id
                )
            )

    # =====================================================
    # POST
    # =====================================================

    if request.method == "POST":

        # =================================================
        # ADMIN
        # =================================================

        if role == "ADMIN":

            # ---------------------------------------------
            # BASIC INFORMATION
            # ---------------------------------------------

            company = request.form.get(
                "company"
            )

            mri_no = request.form.get(
                "mri_no"
            )

            requested_by = request.form.get(
                "requested_by"
            )

            req_date = request.form.get(
                "req_date"
            )

            if company:

                material_request.company = (
                    company
                )

            if mri_no:

                material_request.mri_no = (
                    mri_no
                )

            if requested_by:

                material_request.requested_by = (
                    requested_by
                )

            if req_date:

                try:

                    material_request.req_date = (
                        datetime.strptime(
                            req_date,
                            "%Y-%m-%d"
                        ).date()
                    )

                except ValueError:

                    db.session.rollback()

                    flash(
                        "Invalid Request Date.",
                        "danger"
                    )

                    return redirect(
                        url_for(
                            "edit_material_request",
                            id=id
                        )
                    )

            # ---------------------------------------------
            # LTSC STATUS
            # ---------------------------------------------

            ltsc_status = request.form.get(
                "ltsc_status",
                material_request.ltsc_status
            )

            if ltsc_status in [
                "Submitted",
                "Cancelled"
            ]:

                material_request.ltsc_status = (
                    ltsc_status
                )

            # ---------------------------------------------
            # PIPING STATUS
            # ---------------------------------------------

            piping_status = request.form.get(
                "piping_status",
                material_request.piping_status
            )

            if piping_status in [
                "Pending",
                "In Progress",
                "Completed",
                "Approved",
                "Rejected"
            ]:

                material_request.piping_status = (
                    piping_status
                )

            # ---------------------------------------------
            # PIPING DATE
            # ---------------------------------------------

            piping_date = request.form.get(
                "piping_date"
            )

            if piping_date:

                try:

                    material_request.piping_date = (
                        datetime.strptime(
                            piping_date,
                            "%Y-%m-%d"
                        ).date()
                    )

                except ValueError:

                    db.session.rollback()

                    flash(
                        "Invalid Piping Date.",
                        "danger"
                    )

                    return redirect(
                        url_for(
                            "edit_material_request",
                            id=id
                        )
                    )

            else:

                material_request.piping_date = None

            # ---------------------------------------------
            # CCS-JV
            # ---------------------------------------------

            ccs_jv_confirm = request.form.get(
                "ccs_jv_confirm",
                material_request.ccs_jv_confirm
            )

            if ccs_jv_confirm in [
                "Pending",
                "Confirmed",
                "Rejected"
            ]:

                material_request.ccs_jv_confirm = (
                    ccs_jv_confirm
                )

            # ---------------------------------------------
            # MATERIAL STATUS
            # ---------------------------------------------

            material_status = request.form.get(
                "material_status",
                material_request.material_status
            )

            if material_status in [
                "Pending",
                "Approved",
                "Ordered",
                "In Transit",
                "Received",
                "Rejected"
            ]:

                material_request.material_status = (
                    material_status
                )

            # ---------------------------------------------
            # MATERIAL DATE
            # ---------------------------------------------

            mtrl_date = request.form.get(
                "mtrl_date"
            )

            if mtrl_date:

                try:

                    material_request.mtrl_date = (
                        datetime.strptime(
                            mtrl_date,
                            "%Y-%m-%d"
                        ).date()
                    )

                except ValueError:

                    db.session.rollback()

                    flash(
                        "Invalid Material Date.",
                        "danger"
                    )

                    return redirect(
                        url_for(
                            "edit_material_request",
                            id=id
                        )
                    )

            else:

                material_request.mtrl_date = None

            # ---------------------------------------------
            # REMARKS
            # ---------------------------------------------

            material_request.remarks = (
                request.form.get(
                    "remarks"
                )
            )

        # =================================================
        # REQUESTER
        # =================================================

        elif role == "REQUESTER":

            # ---------------------------------------------
            # BASIC INFORMATION
            # ---------------------------------------------

            company = request.form.get(
                "company"
            )

            mri_no = request.form.get(
                "mri_no"
            )

            requested_by = request.form.get(
                "requested_by"
            )

            req_date = request.form.get(
                "req_date"
            )

            if company:

                material_request.company = (
                    company
                )

            if mri_no:

                material_request.mri_no = (
                    mri_no
                )

            if requested_by:

                material_request.requested_by = (
                    requested_by
                )

            if req_date:

                try:

                    material_request.req_date = (
                        datetime.strptime(
                            req_date,
                            "%Y-%m-%d"
                        ).date()
                    )

                except ValueError:

                    db.session.rollback()

                    flash(
                        "Invalid Request Date.",
                        "danger"
                    )

                    return redirect(
                        url_for(
                            "edit_material_request",
                            id=id
                        )
                    )

            # ---------------------------------------------
            # LTSC STATUS
            # ---------------------------------------------

            ltsc_status = request.form.get(
                "ltsc_status",
                material_request.ltsc_status
            )

            if ltsc_status in [
                "Submitted",
                "Cancelled"
            ]:

                material_request.ltsc_status = (
                    ltsc_status
                )

            # ---------------------------------------------
            # REMARKS
            # ---------------------------------------------

            material_request.remarks = (
                request.form.get(
                    "remarks"
                )
            )

        # =================================================
        # APPROVER PIPING
        # =================================================

        elif role == "APPROVER_PIPING":

            # ---------------------------------------------
            # PIPING STATUS
            # ---------------------------------------------

            piping_status = request.form.get(
                "piping_status",
                material_request.piping_status
            )

            if piping_status in [
                "Pending",
                "In Progress",
                "Completed",
                "Approved",
                "Rejected"
            ]:

                material_request.piping_status = (
                    piping_status
                )

            # ---------------------------------------------
            # PIPING DATE
            # ---------------------------------------------

            piping_date = request.form.get(
                "piping_date"
            )

            if piping_date:

                try:

                    material_request.piping_date = (
                        datetime.strptime(
                            piping_date,
                            "%Y-%m-%d"
                        ).date()
                    )

                except ValueError:

                    db.session.rollback()

                    flash(
                        "Invalid Piping Date.",
                        "danger"
                    )

                    return redirect(
                        url_for(
                            "edit_material_request",
                            id=id
                        )
                    )

            else:

                material_request.piping_date = None

            # ---------------------------------------------
            # CCS-JV
            # ---------------------------------------------

            ccs_jv_confirm = request.form.get(
                "ccs_jv_confirm",
                material_request.ccs_jv_confirm
            )

            if ccs_jv_confirm in [
                "Pending",
                "Confirmed",
                "Rejected"
            ]:

                material_request.ccs_jv_confirm = (
                    ccs_jv_confirm
                )

            # ---------------------------------------------
            # REMARKS
            # ---------------------------------------------

            material_request.remarks = (
                request.form.get(
                    "remarks"
                )
            )

        # =================================================
        # APPROVER MATERIAL
        # =================================================

        elif role == "APPROVER_MATERIAL":

            # ---------------------------------------------
            # MATERIAL STATUS
            # ---------------------------------------------

            material_status = request.form.get(
                "material_status",
                material_request.material_status
            )

            if material_status in [
                "Pending",
                "Approved",
                "Ordered",
                "In Transit",
                "Received",
                "Rejected"
            ]:

                material_request.material_status = (
                    material_status
                )

            # ---------------------------------------------
            # MATERIAL DATE
            # ---------------------------------------------

            mtrl_date = request.form.get(
                "mtrl_date"
            )

            if mtrl_date:

                try:

                    material_request.mtrl_date = (
                        datetime.strptime(
                            mtrl_date,
                            "%Y-%m-%d"
                        ).date()
                    )

                except ValueError:

                    db.session.rollback()

                    flash(
                        "Invalid Material Date.",
                        "danger"
                    )

                    return redirect(
                        url_for(
                            "edit_material_request",
                            id=id
                        )
                    )

            else:

                material_request.mtrl_date = None

            # ---------------------------------------------
            # REMARKS
            # ---------------------------------------------

            material_request.remarks = (
                request.form.get(
                    "remarks"
                )
            )

        # =================================================
        # STATUS VALIDATION
        # =================================================

        # -------------------------------------------------
        # RULE 1
        #
        # Submitted -> Cancelled is allowed ONLY
        # while Piping Status is Pending.
        # -------------------------------------------------

        if (
            original_ltsc_status == "Submitted"
            and material_request.ltsc_status
            == "Cancelled"
            and material_request.piping_status
            != "Pending"
        ):

            db.session.rollback()

            flash(
                "Status cannot be change",
                "danger"
            )

            return redirect(
                url_for(
                    "edit_material_request",
                    id=id
                )
            )

        # -------------------------------------------------
        # RULE 2
        #
        # Material can only be Approved after
        # Piping has been Approved.
        # -------------------------------------------------

        if (
            material_request.material_status
            == "Approved"
            and material_request.piping_status
            != "Approved"
        ):

            db.session.rollback()

            flash(
                "Material status cannot be Approved "
                "until Piping Status is Approved.",
                "danger"
            )

            return redirect(
                url_for(
                    "edit_material_request",
                    id=id
                )
            )

        # =================================================
        # EXCEL FILE
        # =================================================

        excel_file = request.files.get(
            "attachment_excel"
        )

        if (
            excel_file
            and excel_file.filename
        ):

            if not allowed_file(
                excel_file.filename,
                {"xls", "xlsx"}
            ):

                flash(
                    "Invalid Excel file.",
                    "danger"
                )

                return redirect(
                    url_for(
                        "edit_material_request",
                        id=id
                    )
                )

            excel_filename = secure_filename(
                excel_file.filename
            )

            excel_file.save(
                os.path.join(
                    app.config["UPLOAD_FOLDER"],
                    excel_filename
                )
            )

            material_request.attachment_excel = (
                excel_filename
            )

        # =================================================
        # PDF FILE
        # =================================================

        pdf_file = request.files.get(
            "attachment_pdf"
        )

        if (
            pdf_file
            and pdf_file.filename
        ):

            if not allowed_file(
                pdf_file.filename,
                {"pdf"}
            ):

                flash(
                    "Invalid PDF file.",
                    "danger"
                )

                return redirect(
                    url_for(
                        "edit_material_request",
                        id=id
                    )
                )

            pdf_filename = secure_filename(
                pdf_file.filename
            )

            pdf_file.save(
                os.path.join(
                    app.config["UPLOAD_FOLDER"],
                    pdf_filename
                )
            )

            material_request.attachment_pdf = (
                pdf_filename
            )

        # =================================================
        # VOUCHER FILES
        # =================================================

        if role in [
            "ADMIN",
            "APPROVER_MATERIAL"
        ]:

            voucher_files = (
                request.files.getlist(
                    "voucher_files"
                )
            )

            new_voucher_filenames = []

            for voucher_file in voucher_files:

                if not voucher_file:
                    continue

                if not voucher_file.filename:
                    continue

                if not allowed_file(
                    voucher_file.filename,
                    {"pdf"}
                ):

                    flash(
                        "Voucher files must be PDF files.",
                        "danger"
                    )

                    return redirect(
                        url_for(
                            "edit_material_request",
                            id=id
                        )
                    )

                voucher_filename = secure_filename(
                    voucher_file.filename
                )

                voucher_file.save(
                    os.path.join(
                        app.config["UPLOAD_FOLDER"],
                        voucher_filename
                    )
                )

                new_voucher_filenames.append(
                    voucher_filename
                )

            # ---------------------------------------------
            # REPLACE VOUCHERS ONLY WHEN NEW FILES EXIST
            # ---------------------------------------------

            if new_voucher_filenames:

                material_request.voucher_files = (
                    ",".join(
                        new_voucher_filenames
                    )
                )
        # =================================================
        # NOTIFICATIONS
        # =================================================

        # =================================================
        # PIPING STATUS CHANGED
        # ADMIN + PIPING APPROVER
        # =================================================

        if (
            role in [
                "APPROVER_PIPING",
                "ADMIN"
            ]
            and original_piping_status
            != material_request.piping_status
        ):

            # -------------------------------------------------
            # PIPING APPROVED
            # -------------------------------------------------

            if material_request.piping_status == "Approved":

                # Notify requester
                notify_requester(
                    material_request,
                    title="Piping Approved",
                    message=(
                        f"{material_request.mri_no} "
                        f"has been approved by Piping."
                    ),
                    notification_type="APPROVED"
                )

                # Notify ALL Material Approvers
                notify_users_by_role(
                    role="APPROVER_MATERIAL",
                    title="Piping Approved",
                    message=(
                        f"{material_request.mri_no} "
                        f"has passed Piping Approval "
                        f"and is ready for Material "
                        f"processing."
                    ),
                    notification_type="APPROVED",
                    material_request_id=material_request.id
                )

                # Notify ALL Admins
                notify_admins(
                    title="Piping Approved",
                    message=(
                        f"{material_request.mri_no} "
                        f"has been approved by Piping."
                    ),
                    notification_type="APPROVED",
                    material_request_id=material_request.id
                )

            # -------------------------------------------------
            # PIPING REJECTED
            # -------------------------------------------------

            elif material_request.piping_status == "Rejected":

                # Notify requester
                notify_requester(
                    material_request,
                    title="Piping Rejected",
                    message=(
                        f"{material_request.mri_no} "
                        f"has been rejected by Piping."
                    ),
                    notification_type="REJECTED"
                )

                # Notify ALL Admins
                notify_admins(
                    title="Piping Rejected",
                    message=(
                        f"{material_request.mri_no} "
                        f"has been rejected by Piping."
                    ),
                    notification_type="REJECTED",
                    material_request_id=material_request.id
                )


        # =================================================
        # MATERIAL STATUS CHANGED
        # ADMIN + MATERIAL APPROVER
        # =================================================

        if (
            role in [
                "APPROVER_MATERIAL",
                "ADMIN"
            ]
            and original_material_status
            != material_request.material_status
        ):

            # -------------------------------------------------
            # MATERIAL APPROVED
            # -------------------------------------------------

            if material_request.material_status == "Approved":

                # Notify requester
                notify_requester(
                    material_request,
                    title="Material Approved",
                    message=(
                        f"{material_request.mri_no} "
                        f"has been approved by Material."
                    ),
                    notification_type="APPROVED"
                )

                # Notify ALL Admins
                notify_admins(
                    title="Material Approved",
                    message=(
                        f"{material_request.mri_no} "
                        f"has been approved by Material."
                    ),
                    notification_type="APPROVED",
                    material_request_id=material_request.id
                )

            # -------------------------------------------------
            # MATERIAL REJECTED
            # -------------------------------------------------

            elif material_request.material_status == "Rejected":

                # Notify requester
                notify_requester(
                    material_request,
                    title="Material Rejected",
                    message=(
                        f"{material_request.mri_no} "
                        f"has been rejected by Material."
                    ),
                    notification_type="REJECTED"
                )

                # Notify ALL Admins
                notify_admins(
                    title="Material Rejected",
                    message=(
                        f"{material_request.mri_no} "
                        f"has been rejected by Material."
                    ),
                    notification_type="REJECTED",
                    material_request_id=material_request.id
                )
        

        # =================================================
        # SAVE DATABASE CHANGES
        # =================================================

        try:

            db.session.commit()

            flash(
                "Material Request updated successfully.",
                "success"
            )

            return redirect(
                url_for(
                    "view_material_request",
                    id=id
                )
            )

        except Exception as e:

            db.session.rollback()

            flash(
                f"Error updating request: {e}",
                "danger"
            )

            return redirect(
                url_for(
                    "edit_material_request",
                    id=id
                )
            )

    # =====================================================
    # DISPLAY EDIT PAGE
    # =====================================================

    return render_template(
        "material_request_edit.html",
        item=material_request
    )


# =========================================================
# DELETE MATERIAL REQUEST
# =========================================================

@app.route(
    "/material-request/<int:id>/delete",
    methods=["POST"]
)
@login_required
def delete_material_request(id):

    if current_user.role != "ADMIN":

        flash(
            "Only administrators can delete "
            "Material Requests.",
            "danger"
        )

        return redirect(
            url_for("home")
        )

    material_request = (
        MaterialRequest.query
        .get_or_404(id)
    )

    try:

        # -------------------------------------------------
        # DELETE REQUEST
        # -------------------------------------------------

        db.session.delete(
            material_request
        )

        db.session.commit()

        # -------------------------------------------------
        # CHECK REMAINING REQUESTS
        # -------------------------------------------------

        remaining_requests = (
            MaterialRequest.query.count()
        )

        # -------------------------------------------------
        # RESET ID WHEN TABLE IS EMPTY
        # -------------------------------------------------

        if remaining_requests == 0:

            db.session.execute(
                text(
                    "ALTER SEQUENCE "
                    "material_requests_id_seq "
                    "RESTART WITH 1"
                )
            )

            db.session.commit()

        flash(
            "Material Request deleted.",
            "success"
        )

    except Exception as e:

        db.session.rollback()

        flash(
            f"Error deleting request: {e}",
            "danger"
        )

    return redirect(
        url_for("home")
    )


# =========================================================
# DOWNLOAD FILE
# =========================================================

@app.route(
    "/uploads/<filename>"
)
@login_required
def uploaded_file(filename):

    return send_from_directory(
        app.config["UPLOAD_FOLDER"],
        filename,
        as_attachment=True
    )


# =========================================================
# COMPANY REQUESTS
# =========================================================

@app.route(
    "/company/<company>"
)
@login_required
def company_requests(company):

    allowed_companies = [
        "WOONGNAM",
        "DONG-IL",
        "DAEAH"
    ]

    if company not in allowed_companies:

        flash(
            "Invalid company.",
            "danger"
        )

        return redirect(
            url_for("home")
        )

    search = request.args.get(
        "search",
        ""
    ).strip()

    status = request.args.get(
        "status",
        ""
    ).strip()

    query = MaterialRequest.query.filter(
        MaterialRequest.company.ilike(
            company
        )
    )

    if search:

        search_filter = (
            MaterialRequest.mri_no.ilike(
                f"%{search}%"
            )
            |
            MaterialRequest.requested_by.ilike(
                f"%{search}%"
            )
        )

        query = query.filter(
            search_filter
        )

    if status:

        query = query.filter(
            MaterialRequest.material_status
            == status
        )

    requests = query.order_by(
        MaterialRequest.id.desc()
    ).all()

    return render_template(
        "company_requests.html",
        requests=requests,
        company=company,
        search=search,
        status=status
    )


# =========================================================
# COMPANY EXCEL EXPORT
# =========================================================

@app.route(
    "/company/<company>/export/excel"
)
@login_required
def export_company_excel(company):

    allowed_companies = [
        "WOONGNAM",
        "DONG-IL",
        "DAEAH"
    ]

    if company not in allowed_companies:

        flash(
            "Invalid company.",
            "danger"
        )

        return redirect(
            url_for("home")
        )

    search = request.args.get(
        "search",
        ""
    ).strip()

    status = request.args.get(
        "status",
        ""
    ).strip()

    query = MaterialRequest.query.filter(
        MaterialRequest.company.ilike(
            company
        )
    )

    if search:

        search_filter = (
            MaterialRequest.mri_no.ilike(
                f"%{search}%"
            )
            |
            MaterialRequest.requested_by.ilike(
                f"%{search}%"
            )
        )

        query = query.filter(
            search_filter
        )

    if status:

        query = query.filter(
            MaterialRequest.material_status
            == status
        )

    requests = query.order_by(
        MaterialRequest.id.desc()
    ).all()

    workbook = Workbook()

    worksheet = cast(
        Worksheet,
        workbook.active
    )

    worksheet.title = company[:31]

    worksheet["A1"] = (
        f"{company} - MATERIAL REQUEST REPORT"
    )

    worksheet["A1"].font = Font(
        bold=True,
        size=16
    )

    worksheet["A2"] = (
        "Generated: "
        f"{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    )

    worksheet["A3"] = (
        f"Total Records: {len(requests)}"
    )

    worksheet["A4"] = "Company:"
    worksheet["B4"] = company

    worksheet["C4"] = "Search:"
    worksheet["D4"] = (
        search if search else "All"
    )

    worksheet["E4"] = (
        "Material Status:"
    )

    worksheet["F4"] = (
        status if status else "All"
    )

    headers = [
        "ID",
        "Company",
        "MRI No.",
        "Requested By",
        "Request Date",
        "LTSC Status",
        "Piping Status",
        "Piping Date",
        "CCS-JV Confirmation",
        "Material Status",
        "Material Date",
        "Excel File",
        "PDF File",
        "Voucher Files",
        "Remarks",
        "Created At",
        "Updated At"
    ]

    header_row = 6

    for column, header in enumerate(
        headers,
        start=1
    ):

        cell = cast(
            Cell,
            worksheet.cell(
                row=header_row,
                column=column
            )
        )

        cell.value = header

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    row_number = header_row + 1

    for item in requests:

        values = [
            item.id,
            item.company,
            item.mri_no,
            item.requested_by,
            item.req_date,
            item.ltsc_status,
            item.piping_status,
            item.piping_date,
            item.ccs_jv_confirm,
            item.material_status,
            item.mtrl_date,
            item.attachment_excel,
            item.attachment_pdf,
            item.voucher_files,
            item.remarks,
            item.created_at,
            item.updated_at
        ]

        for column, value in enumerate(
            values,
            start=1
        ):

            cell = cast(
                Cell,
                worksheet.cell(
                    row=row_number,
                    column=column
                )
            )

            cell.value = value

            cell.alignment = Alignment(
                vertical="top"
            )

        worksheet.cell(
            row=row_number,
            column=5
        ).number_format = (
            "yyyy-mm-dd"
        )

        worksheet.cell(
            row=row_number,
            column=8
        ).number_format = (
            "yyyy-mm-dd"
        )

        worksheet.cell(
            row=row_number,
            column=11
        ).number_format = (
            "yyyy-mm-dd"
        )

        worksheet.cell(
            row=row_number,
            column=16
        ).number_format = (
            "yyyy-mm-dd hh:mm:ss"
        )

        worksheet.cell(
            row=row_number,
            column=17
        ).number_format = (
            "yyyy-mm-dd hh:mm:ss"
        )

        row_number += 1

    worksheet.freeze_panes = "A7"

    if requests:

        last_row = (
            header_row
            + len(requests)
        )

        worksheet.auto_filter.ref = (
            f"A{header_row}:Q{last_row}"
        )

    widths = {
        "A": 10,
        "B": 20,
        "C": 18,
        "D": 22,
        "E": 15,
        "F": 15,
        "G": 18,
        "H": 15,
        "I": 22,
        "J": 20,
        "K": 15,
        "L": 30,
        "M": 30,
        "N": 40,
        "O": 40,
        "P": 22,
        "Q": 22
    }

    for column, width in widths.items():

        worksheet.column_dimensions[
            column
        ].width = width

    worksheet.row_dimensions[
        header_row
    ].height = 30

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    filename = (
        f"{company}_Material_Request_Report_"
        f"{datetime.now().strftime('%Y-%m-%d')}.xlsx"
    )

    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )


# =========================================================
# USER MANAGEMENT
# =========================================================

@app.route(
    "/users",
    methods=["GET", "POST"]
)
@login_required
def users():

    if current_user.role != "ADMIN":

        flash(
            "Only administrators can manage users.",
            "danger"
        )

        return redirect(
            url_for("home")
        )

    if request.method == "POST":

        username = request.form.get(
            "username",
            ""
        ).strip()

        full_name = request.form.get(
            "full_name",
            ""
        ).strip()

        role = request.form.get(
            "role",
            ""
        ).strip().upper()

        password = request.form.get(
            "password",
            ""
        )

        confirm_password = request.form.get(
            "confirm_password",
            ""
        )

        allowed_roles = [
            "ADMIN",
            "REQUESTER",
            "APPROVER_PIPING",
            "APPROVER_MATERIAL"
        ]

        # -------------------------------------------------
        # VALIDATION
        # -------------------------------------------------

        if not username:

            flash(
                "Username is required.",
                "danger"
            )

            return redirect(
                url_for("users")
            )

        if not full_name:

            flash(
                "Full name is required.",
                "danger"
            )

            return redirect(
                url_for("users")
            )

        if role not in allowed_roles:

            flash(
                "Invalid user role.",
                "danger"
            )

            return redirect(
                url_for("users")
            )

        if not password:

            flash(
                "Password is required.",
                "danger"
            )

            return redirect(
                url_for("users")
            )

        if password != confirm_password:

            flash(
                "Passwords do not match.",
                "danger"
            )

            return redirect(
                url_for("users")
            )

        existing_user = User.query.filter_by(
            username=username
        ).first()

        if existing_user:

            flash(
                "Username already exists.",
                "danger"
            )

            return redirect(
                url_for("users")
            )

        # -------------------------------------------------
        # CREATE USER
        # -------------------------------------------------

        user = User()

        user.username = username

        user.full_name = full_name

        user.role = role

        user.set_password(
            password
        )

        db.session.add(
            user
        )

        db.session.commit()

        flash(
            f"{role} user '{username}' "
            "created successfully.",
            "success"
        )

        return redirect(
            url_for("users")
        )

    all_users = (
        User.query
        .order_by(
            User.id.asc()
        )
        .all()
    )

    return render_template(
        "users.html",
        users=all_users
    )


# =========================================================
# EDIT USER
# =========================================================

@app.route(
    "/users/<int:id>/edit",
    methods=["GET", "POST"]
)
@login_required
def edit_user(id):

    if current_user.role != "ADMIN":

        flash(
            "Only administrators can edit users.",
            "danger"
        )

        return redirect(
            url_for("home")
        )

    user = User.query.get_or_404(
        id
    )

    allowed_roles = [
        "ADMIN",
        "REQUESTER",
        "APPROVER_PIPING",
        "APPROVER_MATERIAL"
    ]

    if request.method == "POST":

        username = request.form.get(
            "username",
            ""
        ).strip()

        full_name = request.form.get(
            "full_name",
            ""
        ).strip()

        role = request.form.get(
            "role",
            ""
        ).strip().upper()

        if not username:

            flash(
                "Username is required.",
                "danger"
            )

            return redirect(
                url_for(
                    "edit_user",
                    id=id
                )
            )

        if not full_name:

            flash(
                "Full name is required.",
                "danger"
            )

            return redirect(
                url_for(
                    "edit_user",
                    id=id
                )
            )

        if role not in allowed_roles:

            flash(
                "Invalid user role.",
                "danger"
            )

            return redirect(
                url_for(
                    "edit_user",
                    id=id
                )
            )

        existing_user = User.query.filter(
            User.username == username,
            User.id != user.id
        ).first()

        if existing_user:

            flash(
                "Username already exists.",
                "danger"
            )

            return redirect(
                url_for(
                    "edit_user",
                    id=id
                )
            )

        user.username = username

        user.full_name = full_name

        user.role = role

        db.session.commit()

        flash(
            f"User '{username}' updated successfully.",
            "success"
        )

        return redirect(
            url_for("users")
        )

    return render_template(
        "edit_user.html",
        user=user,
        allowed_roles=allowed_roles
    )


# =========================================================
# CHANGE USER PASSWORD
# =========================================================

@app.route(
    "/users/<int:id>/password",
    methods=["GET", "POST"]
)
@login_required
def change_user_password(id):

    if current_user.role != "ADMIN":

        flash(
            "Only administrators can change passwords.",
            "danger"
        )

        return redirect(
            url_for("home")
        )

    user = User.query.get_or_404(
        id
    )

    if request.method == "POST":

        password = request.form.get(
            "password",
            ""
        )

        confirm_password = request.form.get(
            "confirm_password",
            ""
        )

        if not password:

            flash(
                "Password is required.",
                "danger"
            )

            return redirect(
                url_for(
                    "change_user_password",
                    id=id
                )
            )

        if password != confirm_password:

            flash(
                "Passwords do not match.",
                "danger"
            )

            return redirect(
                url_for(
                    "change_user_password",
                    id=id
                )
            )

        user.set_password(
            password
        )

        db.session.commit()

        flash(
            f"Password for '{user.username}' "
            "changed successfully.",
            "success"
        )

        return redirect(
            url_for("users")
        )

    return render_template(
        "change_user_password.html",
        user=user
    )


# =========================================================
# TOGGLE USER STATUS
# =========================================================

@app.route(
    "/users/<int:id>/toggle-status",
    methods=["POST"]
)
@login_required
def toggle_user_status(id):

    if current_user.role != "ADMIN":

        flash(
            "Only administrators can change "
            "user status.",
            "danger"
        )

        return redirect(
            url_for("home")
        )

    user = User.query.get_or_404(
        id
    )

    # -----------------------------------------------------
    # PREVENT SELF DEACTIVATION
    # -----------------------------------------------------

    if user.id == current_user.id:

        flash(
            "You cannot deactivate your own account.",
            "danger"
        )

        return redirect(
            url_for("users")
        )

    new_status = (
        not bool(user.is_active)
    )

    # -----------------------------------------------------
    # SQLALCHEMY UPDATE
    # -----------------------------------------------------

    db.session.query(
        User
    ).filter(
        User.id == user.id
    ).update(
        {
            "is_active": new_status
        }
    )

    db.session.commit()

    if new_status:

        flash(
            f"User '{user.username}' "
            "has been activated.",
            "success"
        )

    else:

        flash(
            f"User '{user.username}' "
            "has been deactivated.",
            "success"
        )

    return redirect(
        url_for("users")
    )


# =========================================================
# DELETE USER
# =========================================================

@app.route(
    "/users/<int:id>/delete",
    methods=["POST"]
)
@login_required
def delete_user(id):

    if current_user.role != "ADMIN":

        flash(
            "Only administrators can delete users.",
            "danger"
        )

        return redirect(
            url_for("home")
        )

    user = User.query.get_or_404(
        id
    )

    # -----------------------------------------------------
    # PREVENT SELF DELETE
    # -----------------------------------------------------

    if user.id == current_user.id:

        flash(
            "You cannot delete your own account.",
            "danger"
        )

        return redirect(
            url_for("users")
        )

    # -----------------------------------------------------
    # CHECK APPROVAL HISTORY
    # -----------------------------------------------------

    history_count = (
        ApprovalHistory.query
        .filter_by(
            user_id=user.id
        )
        .count()
    )

    if history_count > 0:

        flash(
            "This user cannot be deleted because "
            "they have approval history. "
            "Deactivate the account instead.",
            "danger"
        )

        return redirect(
            url_for("users")
        )

    username = user.username

    try:

        db.session.delete(
            user
        )

        db.session.commit()

        flash(
            f"User '{username}' "
            "deleted successfully.",
            "success"
        )

    except Exception as e:

        db.session.rollback()

        flash(
            f"Error deleting user: {e}",
            "danger"
        )

    return redirect(
        url_for("users")
    )


# =========================================================
# NOTIFICATIONS PAGE
# =========================================================

@app.route(
    "/notifications"
)
@login_required
def notifications():

    user_notifications = (
        Notification.query
        .filter_by(
            user_id=current_user.id
        )
        .order_by(
            Notification.created_at.desc()
        )
        .all()
    )

    return render_template(
        "notifications.html",
        notifications=user_notifications
    )


# =========================================================
# MARK ONE NOTIFICATION AS READ
# =========================================================

@app.route(
    "/notifications/<int:id>/read",
    methods=["POST"]
)
@login_required
def mark_notification_read(id):

    notification = (
        Notification.query
        .filter_by(
            id=id,
            user_id=current_user.id
        )
        .first_or_404()
    )

    notification.is_read = True

    db.session.commit()

    # -----------------------------------------------------
    # OPEN RELATED MATERIAL REQUEST
    # -----------------------------------------------------

    if notification.material_request_id:

        return redirect(
            url_for(
                "view_material_request",
                id=notification.material_request_id
            )
        )

    return redirect(
        url_for("notifications")
    )


# =========================================================
# MARK ALL NOTIFICATIONS AS READ
# =========================================================

@app.route(
    "/notifications/read-all",
    methods=["POST"]
)
@login_required
def mark_all_notifications_read():

    Notification.query.filter_by(
        user_id=current_user.id,
        is_read=False
    ).update(
        {
            "is_read": True
        }
    )

    db.session.commit()

    flash(
        "All notifications marked as read.",
        "success"
    )

    return redirect(
        url_for("notifications")
    )


# =========================================================
# GLOBAL NOTIFICATION COUNT
# =========================================================

@app.context_processor
def inject_notifications():

    if not current_user.is_authenticated:

        return {
            "unread_notifications": 0
        }

    unread_count = (
        Notification.query
        .filter_by(
            user_id=current_user.id,
            is_read=False
        )
        .count()
    )

    return {
        "unread_notifications": unread_count
    }
    
    
@app.route("/notifications/<int:id>/delete", methods=["POST"])
@login_required
def delete_notification(id):
    notification = Notification.query.filter_by(
        id=id,
        user_id=current_user.id
    ).first_or_404()

    db.session.delete(notification)
    db.session.commit()

    flash("Notification deleted.", "success")
    return redirect(url_for("notifications"))


@app.route("/notifications/delete-all", methods=["POST"])
@login_required
def delete_all_notifications():

    Notification.query.filter_by(
        user_id=current_user.id
    ).delete()

    db.session.commit()

    flash("All notifications deleted.", "success")
    return redirect(url_for("notifications"))



# =========================================================
# START APPLICATION
# =========================================================

if __name__ == "__main__":

    with app.app_context():

        db.create_all()

    app.run(
        debug=True,
        host="127.0.0.1",
        port=5000
    )